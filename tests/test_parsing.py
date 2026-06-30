"""Test d'intégration de la brique 4b (parsing + OCR), en conteneur.

Génère à la volée un Markdown, un DOCX et une image (texte rendu), puis les
indexe via le pipeline complet (parsing -> chunks -> tagging -> dense+sparse).
Vérifie l'extraction native (fast), l'OCR (fra+eng) et la déduplication.

Prérequis : Ollama (bge-m3) + Qdrant accessibles ; Tesseract installé (conteneur).
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

APP_DIR = Path(__file__).resolve().parent.parent / "app"
sys.path.insert(0, str(APP_DIR))

from core import bases, ingest, vectorstore  # noqa: E402

_failures: list[str] = []


def check(label: str, cond: bool) -> None:
    print(f"[{'OK  ' if cond else 'ECHEC'}] {label}")
    if not cond:
        _failures.append(label)


def _make_markdown(d: Path) -> Path:
    p = d / "doctrine.md"
    p.write_text(
        "# Frégates de la Marine nationale\n\n"
        "La frégate multi-missions FREMM assure la lutte anti-sous-marine.\n\n"
        "## Capteurs\n\n"
        "Le sonar remorqué détecte les sous-marins à grande distance.\n",
        encoding="utf-8",
    )
    return p


def _make_docx(d: Path) -> Path:
    import docx
    doc = docx.Document()
    doc.add_heading("Porte-avions Charles de Gaulle", level=1)
    doc.add_paragraph("Le porte-avions assure la projection de puissance aérienne "
                      "avec son groupe aérien embarqué de Rafale Marine.")
    p = d / "porte-avions.docx"
    doc.save(str(p))
    return p


def _make_image(d: Path) -> Path:
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 200), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 30)
    draw.text((20, 30), "Procedure de securite", fill="black", font=font)
    draw.text((20, 90), "Naval safety document", fill="black", font=font)
    p = d / "scan.png"
    img.save(str(p))
    return p


def main() -> int:
    bid = None
    try:
        b = bases.create_base("Test Parsing", "4b", parse_strategy="fast")
        bid = b["id"]

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)

            # --- Markdown (extraction native) ---
            rep = ingest.ingest_file(bid, str(_make_markdown(d)), strategy="fast")
            check("Markdown : chunks indexés", rep.get("indexed", 0) > 0)

            # --- DOCX (extraction native) ---
            rep = ingest.ingest_file(bid, str(_make_docx(d)), strategy="fast")
            check("DOCX : chunks indexés", rep.get("indexed", 0) > 0)

            # --- Image (OCR forcé fra+eng) ---
            rep_img = ingest.ingest_file(bid, str(_make_image(d)), strategy="fast")
            check("Image : OCR -> chunk indexé", rep_img.get("indexed", 0) >= 1)

            total = vectorstore.count_points(bid)
            check("comptage total > 0", total > 0)
            print("       points indexés:", total)

            # --- Déduplication : ré-ingestion du Markdown ignorée ---
            rep_dup = ingest.ingest_file(bid, str(_make_markdown(d)), strategy="fast")
            check("déduplication : Markdown déjà indexé", rep_dup.get("skipped") is True)

            # --- Recherche : le contenu OCR est retrouvable ---
            from core import llm, sparse
            qvec = llm.embed_query("procédure de sécurité navale")
            qidx, qval = sparse.encode_query("procédure de sécurité navale")
            res = vectorstore.hybrid_search(bid, dense_vec=qvec, sparse_indices=qidx,
                                            sparse_values=qval, limit=5)
            blob = " ".join(r["payload"]["text"].lower() for r in res)
            check("recherche : texte OCR retrouvé ('securite')", "securite" in blob or "sécurité" in blob)
            check("recherche : métadonnée fichier présente",
                  any(r["payload"].get("file") for r in res))

    finally:
        if bid:
            try:
                bases.delete_base(bid, remove_documents=True)
            except Exception:
                pass

    print()
    if _failures:
        print(f"RESULTAT : {len(_failures)} echec(s).")
        return 1
    print("RESULTAT : tous les tests passent.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
